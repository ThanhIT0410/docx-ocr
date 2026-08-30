"""End-to-end smoke test for the export pipeline: `LayoutExportService`
(mode='layout') calls `LayoutReconstructorV2` unconditionally (no more
rule-based/ML switch), and `PlainTextExportService` (mode='plain') skips
reconstruction entirely — see `export_service.py`'s module docstring."""
from __future__ import annotations

import pytest
from docx import Document

from app.schemas import DocumentLayout, ExportPageInput, OcrPageResult
from app.services.export_service import create_export_service


def _page(order: int) -> ExportPageInput:
    layouts = [
        DocumentLayout(bbox=[100, 80, 1200, 160], category="Title", text="De thi mau"),
        DocumentLayout(bbox=[100, 200, 1200, 900], category="Text", text="Cau 1: 2 * 3 * 4 = ?"),
    ]
    ocr = OcrPageResult(origin_width=1654, origin_height=2338, input_width=1654, input_height=2338, layouts=layouts)
    return ExportPageInput(
        pageId=f"page-{order}", order=order, ocrText=ocr, originalImageUrl=f"https://example.test/{order}.jpg"
    )


def test_export_produces_a_valid_docx_via_v2():
    result = create_export_service("De thi thu").export([_page(1), _page(2)])
    assert result.filename == "De thi thu.docx"
    assert result.media_type.endswith("wordprocessingml.document")

    from io import BytesIO

    doc = Document(BytesIO(result.content))
    # 2 pages, each starting its own w:sectPr — this is what
    # start_new_page=not self._is_first_page (document_renderer.py) produces.
    assert len(doc.sections) >= 2


def test_plain_mode_skips_reconstruction_and_keeps_a_real_table():
    """mode='plain' must not run PageFormatNormalizer/LayoutReconstructorV2
    at all — an aspect ratio that would 422 in layout mode should export
    fine here — and a Table block must still become an actual docx table,
    not literal HTML text."""
    layouts = [
        DocumentLayout(bbox=[10, 10, 50, 30], category="Title", text="**Bold** title"),
        DocumentLayout(bbox=[10, 40, 50, 60], category="Text", text="Some body text"),
        DocumentLayout(
            bbox=[10, 70, 50, 90], category="Table",
            text="<table><tr><td>A</td><td>B</td></tr></table>",
        ),
    ]
    # Wildly non-standard aspect ratio — would raise UnsupportedPageAspectRatioError in layout mode.
    ocr = OcrPageResult(origin_width=50, origin_height=9000, input_width=50, input_height=9000, layouts=layouts)
    page = ExportPageInput(pageId="p1", order=1, ocrText=ocr, originalImageUrl="https://example.test/1.jpg")

    result = create_export_service("De thi thu", "plain").export([page])

    from io import BytesIO

    doc = Document(BytesIO(result.content))
    tables = doc.tables
    assert len(tables) == 1
    assert tables[0].cell(0, 0).text == "A"
    assert tables[0].cell(0, 1).text == "B"
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Some body text" in body_text
    # "Plain" means no bold *run* applied, but the "**" markdown-lite
    # delimiters are still markup, not content — must not leak into the text.
    assert "Bold title" in body_text
    assert "**" not in body_text


def test_pdf_mode_reassembles_original_images_ignoring_ocr_text(tmp_path):
    """mode='pdf' must not touch ocrText at all (a page with none must still
    export fine) — it just re-fetches each page's original image
    (`originalImageUrl`, here a local `file://` URL standing in for a
    Supabase signed URL) and stacks them into one PDF, in page order, sized
    to each image's own physical dimensions."""
    from PIL import Image

    from app.services.export_service import RENDER_DPI

    path1 = tmp_path / "p1.png"
    Image.new("RGB", (RENDER_DPI * 2, RENDER_DPI * 3), color="white").save(path1)  # 2in x 3in
    path2 = tmp_path / "p2.png"
    Image.new("RGB", (RENDER_DPI * 4, RENDER_DPI * 1), color="black").save(path2)  # 4in x 1in

    # Deliberately out of order, to check `export()` still sorts by `order`.
    page2 = ExportPageInput(pageId="p2", order=2, ocrText=None, originalImageUrl=path2.as_uri())
    page1 = ExportPageInput(pageId="p1", order=1, ocrText=None, originalImageUrl=path1.as_uri())

    result = create_export_service("De thi anh", "pdf").export([page2, page1])
    assert result.filename == "De thi anh.pdf"
    assert result.media_type == "application/pdf"

    import fitz

    doc = fitz.open(stream=result.content, filetype="pdf")
    assert doc.page_count == 2
    assert doc[0].rect.width == pytest.approx(2 * 72, abs=1)
    assert doc[0].rect.height == pytest.approx(3 * 72, abs=1)
    assert doc[1].rect.width == pytest.approx(4 * 72, abs=1)
    assert doc[1].rect.height == pytest.approx(1 * 72, abs=1)
    doc.close()


def test_pdf_mode_rejects_a_page_with_no_original_image():
    page = ExportPageInput(pageId="p1", order=1, ocrText=None, originalImageUrl="")
    with pytest.raises(ValueError, match="không có ảnh gốc"):
        create_export_service("De thi", "pdf").export([page])


def test_drop_stamp_pictures_removes_only_overlapping_pictures():
    """A `Picture` block overlapping other content is treated as a
    stamp/watermark and dropped before layout reconstruction ever sees it;
    a `Picture` that doesn't overlap anything is real content and kept."""
    from app.services.export_service import LayoutExportService

    stamp = DocumentLayout(bbox=[100, 100, 400, 400], category="Picture", text="")
    text_under_stamp = DocumentLayout(bbox=[150, 150, 350, 350], category="Text", text="real content")
    touching_not_overlapping = DocumentLayout(bbox=[400, 100, 600, 400], category="Text", text="beside it")
    standalone_picture = DocumentLayout(bbox=[600, 600, 800, 800], category="Picture", text="")

    service = LayoutExportService("De thi")
    kept = service._drop_stamp_pictures([stamp, text_under_stamp, touching_not_overlapping, standalone_picture])

    assert stamp not in kept
    assert text_under_stamp in kept
    assert touching_not_overlapping in kept
    assert standalone_picture in kept
